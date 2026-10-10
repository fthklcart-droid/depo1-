"""Metinden belge dosyası üretir: md, html, docx, xlsx, csv. docx/xlsx için python-docx / openpyxl gerekir."""
import html
import re
import time

from core import store

TOOL = {
    "name": "belge_uret",
    "description": "Verilen içerikten dosya oluşturur. bicim: md | html | docx | xlsx | csv. docx'te '#' başlık, '-' madde işareti; xlsx/csv'de her satır bir satır, hücreler '|' ile ayrılır (ilk satır başlık).",
    "parameters": {"type": "object", "properties": {
        "bicim": {"type": "string", "enum": ["md", "html", "docx", "xlsx", "csv"]}, "baslik": {"type": "string"},
        "icerik": {"type": "string"}}, "required": ["bicim", "baslik", "icerik"]},
}


def _slug(t):
    return re.sub(r"[^a-z0-9]+", "-", t.lower().translate(str.maketrans("ığüşöçİ", "igusoci"))).strip("-") or "belge"


def _rows(text):
    return [[c.strip() for c in line.split("|")] for line in text.splitlines() if line.strip()]


def to_html(title, text):
    out, in_ul = [], False
    for line in text.splitlines():
        m = re.match(r"^(#{1,3})\s+(.*)", line)
        li = re.match(r"^\s*[-*]\s+(.*)", line)
        if in_ul and not li:
            out.append("</ul>"); in_ul = False
        if m:
            out.append(f"<h{len(m.group(1))}>{html.escape(m.group(2))}</h{len(m.group(1))}>")
        elif li:
            if not in_ul:
                out.append("<ul>"); in_ul = True
            out.append(f"<li>{html.escape(li.group(1))}</li>")
        elif line.strip():
            out.append(f"<p>{html.escape(line)}</p>")
    if in_ul:
        out.append("</ul>")
    return ("<!doctype html><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>"
            f"<title>{html.escape(title)}</title><body style='font-family:system-ui;max-width:760px;margin:2rem auto;padding:0 1rem;line-height:1.6'>"
            f"<h1>{html.escape(title)}</h1>" + "\n".join(out))


def run(p, ctx):
    fmt, title, text = p["bicim"], p["baslik"], p["icerik"]
    name = f"{_slug(title)}-{time.strftime('%Y%m%d-%H%M%S')}.{fmt}"
    d = store.DIR / "outputs"
    d.mkdir(parents=True, exist_ok=True)
    path = d / name
    if fmt in ("md", "csv"):
        path.write_text(text if fmt == "md" else "\n".join(",".join('"' + c.replace('"', '""') + '"' for c in r) for r in _rows(text)), "utf-8-sig" if fmt == "csv" else "utf-8")
    elif fmt == "html":
        path.write_text(to_html(title, text), "utf-8")
    elif fmt == "docx":
        try:
            from docx import Document
        except ImportError:
            return "docx için python-docx gerekli: pip install python-docx"
        doc = Document()
        doc.add_heading(title, 0)
        for line in text.splitlines():
            m, li = re.match(r"^(#{1,3})\s+(.*)", line), re.match(r"^\s*[-*]\s+(.*)", line)
            if m:
                doc.add_heading(m.group(2), len(m.group(1)))
            elif li:
                doc.add_paragraph(li.group(1), style="List Bullet")
            elif line.strip():
                doc.add_paragraph(line)
        doc.save(path)
    else:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
        except ImportError:
            return "xlsx için openpyxl gerekli: pip install openpyxl"
        wb = Workbook(); ws = wb.active; ws.title = title[:31]
        for r in _rows(text):
            ws.append([float(c.replace(",", ".")) if re.fullmatch(r"-?\d+([.,]\d+)?", c) else c for c in r])
        for c in ws[1]:
            c.font = Font(bold=True)
        wb.save(path)
    ctx.undo.push(f"{fmt} dosyası", lambda: (path.unlink(missing_ok=True), "Dosya silindi.")[1])
    return f"Oluşturuldu: {path} ({path.stat().st_size} bayt)"
